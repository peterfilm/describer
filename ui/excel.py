from openpyxl import Workbook
from openpyxl.drawing.image import Image as OpenPyXLImage
from openpyxl.styles import Alignment, Font
from PIL import Image as PILImage
import os
from PyQt5.QtWidgets import QFileDialog
from datetime import datetime
from utils import load_api_keys, load_key_to_api

# Настройки
thumbnail_width = 200  # Ширина миниатюры
margin = 14  # Запас в пикселях для предотвращения перекрытия
col_c_width = (thumbnail_width + margin) / 7.5
col_d_multiplier = 2.5  # Коэффициент увеличения ширины колонки с описанием

class Excel(object):
    '''
    Для сохранения списка в excel документ
    '''
    def __init__(self, ui):
        self.ui = ui
        self.conf = load_api_keys()
        self.last_path = self.conf['LAST_SAVE_PATH']
        self.ui.action_excel.triggered.connect(self.save_excel)
        
        self.ui.lang_changed.connect(self.retranslateUi)
    
    def retranslateUi(self):
        '''перевод'''
        self.suc_saved = self.ui._translate('Describer', 'Файл {index} успешно создан!')
        self.save_excel_text = self.ui._translate('Describer', 'Сохранить файл Excel')
        self.type_doc = self.ui._translate('Describer', 'Excel Document *.xlsx')
        self.fail_save = self.ui._translate('Describer', 'Не удалось сохранить файл')
        self.title = self.ui._translate('Describer', "AI описание картинок")
        
        self.states = {
            'success': self.ui._translate('Describer', 'Файл успешно создан'),
            'fail': self.ui._translate('Describer', 'Не удалось создать файл'),
            'date_gen': self.ui._translate('Describer', 'Дата генерации'),
        }
        
    
    def save_excel(self):
        try:
            file_path = QFileDialog.getSaveFileName(self.ui, self.save_excel_text, self.last_path, self.type_doc)[0]
            file_name = os.path.basename(file_path)
            data = [[i.number, os.path.abspath(i.path), i.get_description()] for i in self.ui.photos.photos]
            self.create_excel(data, file_path)
            self.ui.timer.on_clicked(self.suc_saved.format(index=file_name), 5000)
        except Exception as e:
            self.ui.timer.on_clicked(self.fail_save, 5000)
            print(e)


    # Функция для создания миниатюры
    @staticmethod
    def create_thumbnail(image_path, target_width=200):
        with PILImage.open(image_path) as img:
            # Пропорциональное уменьшение
            aspect_ratio = img.height / img.width
            new_width = target_width
            new_height = int(target_width * aspect_ratio)
            img = img.resize((new_width, new_height))
            
            # Сохранение миниатюры
            thumbnail_path = f"{os.path.splitext(image_path)[0]}_thumbnail.png"
            img.save(thumbnail_path, "PNG")
            return thumbnail_path, new_height
        
    
    def create_excel(self, data, file_path):
        try:
            # Создание Excel-файла
            wb = Workbook()
            ws = wb.active
            ws.title = self.title 

            # Заголовок на всю ширину
            header = self.title
            ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=4)
            ws["A1"] = header
            ws["A1"].alignment = Alignment(horizontal="center", vertical="center")
            ws["A1"].font = Font(size=14, bold=True)

            # Установка ширины колонок
            ws.column_dimensions["A"].width = 10
            ws.column_dimensions["B"].width = 30
            ws.column_dimensions["C"].width = col_c_width
            ws.column_dimensions["D"].width = col_c_width * col_d_multiplier * 2

            # Список для хранения путей к миниатюрам
            thumbnails_to_delete = []

            # Заполнение данных
            row = 2
            for index, image_path, description in data:
                # Индекс
                cell_a = ws.cell(row=row, column=1, value=index)
                cell_a.alignment = Alignment(horizontal="center", vertical="center")
                
                # Название файла
                filename = os.path.basename(image_path)
                cell_b = ws.cell(row=row, column=2, value=filename)
                cell_b.alignment = Alignment(horizontal="center", vertical="center")
                
                # Миниатюра
                thumbnail_path, thumbnail_height = self.create_thumbnail(image_path, target_width=thumbnail_width)
                thumbnails_to_delete.append(thumbnail_path)  # Добавляем путь к миниатюре в список
                img = OpenPyXLImage(thumbnail_path)
                img.anchor = f"C{row}"
                ws.add_image(img)
                
                # Установка высоты строки
                ws.row_dimensions[row].height = thumbnail_height * 0.75  # Excel использует специальные единицы для высоты строки
                
                # Описание
                cell_d = ws.cell(row=row, column=4, value=description)
                cell_d.alignment = Alignment(wrap_text=True, vertical="top")
                
                row += 1

            # Добавление строки с датой
            date_row = row + 1
            ws.merge_cells(start_row=date_row, start_column=1, end_row=date_row, end_column=4)
            ws[f"A{date_row}"] = f"{self.states['date_gen']}: {datetime.today().strftime('%Y-%m-%d')}"
            ws[f"A{date_row}"].alignment = Alignment(horizontal="left", vertical="center")
            ws.row_dimensions[date_row].height = 25  # Высота строки с датой

            # Сохранение файла
            wb.save(file_path)
            self.ui.timer.on_clicked(self.states['success'])

            # Удаление временных миниатюр
            for thumbnail_path in thumbnails_to_delete:
                if os.path.exists(thumbnail_path):
                    os.remove(thumbnail_path)

        except Exception as e:
            print(e)
            self.ui.timer.on_clicked(self.states['fail'])